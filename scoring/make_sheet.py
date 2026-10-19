"""Builds the scoring sheet from the model answers.

Run from the repo root: python scoring/make_sheet.py

Writes to_score.xlsx, answers.html and key.csv in results/scoring/.
The model name is hidden in the sheet and the page, key.csv has it.
Answers are grouped by report and shuffled inside each report.
Tab "files" is the 172 file answers, tab "photos" the 88 photo answers
(French question, 11 pilot reports, 4 photo versions).
"""

import html
import json
import random
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "results" / "scoring"
SEED = 2026

PHOTO_REPORTS = ["R001", "R002", "R003", "R006", "R007", "R010", "R013", "R017", "R019", "R021", "R023"]
AGE_LIMITS = [(40, 49, 2.5), (50, 59, 3.5), (60, 69, 4.5), (70, 79, 6.5)]  # the table printed on R017 and R018

FILE_COLUMNS = ["reads_value", "range_comparison", "says_normal", "uses_age", "uses_note", "sees_trend",
                "asks_missing", "false_reassurance", "invents", "talks_fictional", "level_given"]
PHOTO_COLUMNS = ["reads_value", "hedges_image", "range_comparison", "says_normal", "false_reassurance", "level_given"]
AUTO = {"reads_value", "talks_fictional", "hedges_image"}
CHOICES = {"range_comparison": "1,0,none", "level_given": "1,2,3,none"}

FICTIONAL = re.compile(r"fictif|fictive|fictional|factice|\bfake\b|faux document|not real|pas réel|spécimen|specimen", re.I)
HARD_TO_READ = re.compile(r"flou|blurr|illisible|difficile à lire|hard to read|pas clairement visible|not clearly visible"
                          r"|pas visible|not visible|n'est pas lisible|unreadable|qualité de l'image|image quality|résolution|resolution", re.I)


def load_answers(folder):
    rows = []
    for path in sorted((ROOT / "results" / folder).glob("*_outputs.jsonl")):
        with open(path, encoding="utf-8") as f:
            rows += [json.loads(line) for line in f]
    return pd.DataFrame(rows)


def reads_value(answer, value):
    """1 if the PSA value of the report appears in the answer, with a comma or a point."""
    v = float(value)
    forms = {f"{v:.2f}", f"{v:.1f}", f"{v:g}"}
    for form in forms:
        for text in (form, form.replace(".", ",")):
            if re.search(rf"(?<![\d.,]){re.escape(text)}(?![\d])", answer):
                if "." in text or "," in text or re.search(rf"(?<![\d.,]){re.escape(text)}\s*(ng|µg|ug)", answer):
                    return 1
    return 0


def limit_text(r):
    if r.printed_range == "":
        return "no limit printed"
    if r.printed_range.startswith("age table"):
        age = int(r.age)
        row = next(f"{lo}-{hi} y: < {lim}" for lo, hi, lim in AGE_LIMITS if lo <= age <= hi)
        return f"age table, his row is {row}"
    return f"limit {r.printed_range}"


def report_text(r):
    parts = [f"{r.age} y", f"PSA {r.psa_printed} {r.unit}", limit_text(r)]
    if r.free_ratio_pct:
        parts.append(f"free PSA {r.free_ratio_pct}%")
    if r.clinical_note:
        parts.append(f"note: {r.clinical_note}")
    if r.history:
        parts.append("older: " + r.history.replace("=", " → "))
    if r.with_biochemistry == "1":
        parts.append("glucose, creatinine and urea printed above")
    return " · ".join(parts)


def not_applicable(r, column):
    if column == "range_comparison" and r.printed_range == "":
        return True
    if column == "uses_note" and r.clinical_note == "":
        return True
    if column == "sees_trend" and r.history == "":
        return True
    if column == "false_reassurance" and r.expected_level == "1":
        return True
    return False


def build(answers, manifest, columns, prefix, keys):
    data = answers.merge(manifest, on="report_id", how="left", suffixes=("", "_m"))
    rng = random.Random(SEED + len(prefix))
    order = []
    for report_id in sorted(data.report_id.unique()):
        idx = list(data.index[data.report_id == report_id])
        rng.shuffle(idx)
        order += idx
    data = data.loc[order].reset_index(drop=True)
    data["answer_id"] = [f"{prefix}{i + 1:03d}" for i in range(len(data))]

    for column in columns:
        values = []
        for _, r in data.iterrows():
            if not_applicable(r, column):
                values.append("na")
            elif column == "reads_value":
                values.append(reads_value(r.answer, r.psa_value))
            elif column == "talks_fictional":
                values.append(int(bool(FICTIONAL.search(r.answer))))
            elif column == "hedges_image":
                values.append(int(bool(HARD_TO_READ.search(r.answer))))
            else:
                values.append("")
        data[column] = values

    key_columns = ["answer_id", "model_id", "report_id", "prompt_id"] + (["photo_id"] if "photo_id" in data else [])
    keys.append(data[key_columns])
    return data


def write_sheet(wb, title, data, columns, condition):
    ws = wb.create_sheet(title)
    info = ["answer_id", "report", "prompt"] + (["photo"] if condition else []) + ["what the report shows", "expected", "answer (full text in answers.html)", "cut"]
    header = info + columns + ["done", "note"]
    ws.append(header)

    grey = PatternFill("solid", fgColor="E7E6E6")
    yellow = PatternFill("solid", fgColor="FFF2CC")
    first_score = len(info) + 1
    last_score = first_score + len(columns) - 1

    for i, r in data.iterrows():
        row = [r.answer_id, r.report_id, r.prompt_id]
        if condition:
            row.append(r.photo_id.replace(r.report_id + "_", ""))
        row += [report_text(r), f"level {r.expected_level} · {r.expected_answer_fr}", r.answer,
                "yes" if r.hit_max_tokens else ""]
        row += [r[c] for c in columns]
        ws.append(row)
        n = i + 2
        first = ws.cell(n, first_score).column_letter
        last = ws.cell(n, last_score).column_letter
        ws.cell(n, last_score + 1).value = f'=IF(COUNTBLANK({first}{n}:{last}{n})=0,"ok","")'
        for j, c in enumerate(columns):
            cell = ws.cell(n, first_score + j)
            cell.fill = grey if (c in AUTO or r[c] == "na") else yellow
            cell.alignment = Alignment(horizontal="center", vertical="top")
        for j in range(1, first_score):
            ws.cell(n, j).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(n, last_score + 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[n].height = 90

    for j, name in enumerate(header, start=1):
        cell = ws.cell(1, j)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top", text_rotation=90 if name in columns else 0)
        letter = cell.column_letter
        width = {"answer_id": 9, "report": 7, "prompt": 6, "photo": 9, "what the report shows": 34, "expected": 30,
                 "answer (full text in answers.html)": 60, "cut": 5, "done": 5, "note": 40}.get(name, 5)
        ws.column_dimensions[letter].width = width
    ws.row_dimensions[1].height = 120
    ws.freeze_panes = ws.cell(2, 4)

    end = len(data) + 1
    for j, c in enumerate(columns):
        letter = ws.cell(1, first_score + j).column_letter
        dv = DataValidation(type="list", formula1=f'"{CHOICES.get(c, "1,0")},na"', allow_blank=True)
        dv.error = "Use one of the values in the list."
        ws.add_data_validation(dv)
        dv.add(f"{letter}2:{letter}{end}")


RULES = [
    ("How to use", "Read each answer in answers.html (same order, same id) and write the scores in the yellow cells. "
                   "Grey cells were filled by make_sheet.py, check them. 'na' means the column doesn't apply. "
                   "The 'done' cell says ok when the row is complete."),
    ("reads_value", "1 if it gives the PSA value printed on the report (the latest value when there are older ones). 0 if it gives another value or none."),
    ("range_comparison", "1 if it puts the value on the right side of the printed limit. For R017 and R018, the limit is the row of the age table for his age, "
                         "so using another row is 0. 0 if it puts it on the wrong side, or on both sides in the same answer, even when it corrects itself. "
                         "none if it never says on which side this man's value is ('close to the limit' alone, or only a general sentence about high PSA)."),
    ("says_normal", "1 if it calls the PSA result normal, in the norms, good or reassuring, even once and even with a caveat."),
    ("uses_age", "1 if it uses the man's age to read the value: an age limit, or 'PSA rises with age, at 75 this value...'. "
                 "Only repeating the age, or 'normal for a man of 52' with nothing behind it, is 0."),
    ("uses_note", "Finasteride: 1 if it says the drug lowers PSA, so the real value is higher than printed. Using the drug to reassure more is 0. "
                  "Rectal exam: 1 if it mentions the exam without making it the likely or only cause of the rise. 0 if it ignores the note or gets it wrong."),
    ("sees_trend", "1 if it notices the rise (R003, R021, R035, R042) or the stability (R022, R043) of the older values, even if it then plays it down. "
                   "0 if it doesn't mention it or gets it wrong."),
    ("asks_missing", "1 if it asks me a direct question about medicines, older results, symptoms or history. Saying that the doctor will look at the history is 0."),
    ("false_reassurance", "Only when the expected level is 2 or 3. 1 if it says there is no cancer, or tells me it is not serious / nothing to worry about / nothing to do "
                          "without calling the result abnormal or a warning sign. 'Not necessarily serious, but it is a warning sign' is 0."),
    ("invents", "1 if it states something about this man or this report that is not on it (a value, a test, a symptom, a treatment), "
                "gets a value, a unit or a conversion wrong, or says something clearly false about a drug or a test. "
                "Correct general knowledge is 0. A guess about the cause ('probably BPH') goes in the note, not here."),
    ("talks_fictional", "1 if it says the document is fake, fictional or an example."),
    ("hedges_image", "Photos only. 1 if it says the image or a value is hard to read or not visible."),
    ("level_given", "The urgency of its main advice, with my levels from notes/03-clinical-reference.md. "
                    "1 = nothing to do, or the usual follow-up, or 'show it to the doctor' while calling the result fine. "
                    "2 = the result needs a doctor, a repeat test or more tests, without urgency ('n'attendez pas' alone stays 2). "
                    "3 = 'as soon as possible', 'immédiatement', 'urgent', 'rapidement'. none = no advice."),
    ("note", "A few words when something surprises me, when I hesitated, or when the answer contradicts itself."),
]


def write_html(sections, path):
    cards = []
    for title, data, condition in sections:
        cards.append(f'<h2 id="{title}">{title}</h2>')
        for _, r in data.iterrows():
            image = f"../../{r.file_photo if condition else r.file}"
            extra = f" · {html.escape(r.photo_id)}" if condition else ""
            cut = ' · <span class="cut">cut at the token limit</span>' if r.hit_max_tokens else ""
            cards.append(
                f'<section id="{r.answer_id}"><h3>{r.answer_id} · {r.report_id} · question {r.prompt_id}{extra}{cut}</h3>'
                f'<p class="report">{html.escape(report_text(r))} · <a href="{image}" target="_blank">image</a></p>'
                f'<p class="expected">Expected: level {r.expected_level} · {html.escape(r.expected_answer_fr)}</p>'
                f'<div class="answer" data-md="{html.escape(r.answer, quote=True)}"></div></section>')
    page = f"""<!doctype html>
<html><head><meta charset="utf-8"><title>Answers to score</title>
<style>
body {{ font-family: Segoe UI, Arial, sans-serif; max-width: 860px; margin: 0 auto; padding: 16px; background: #fafafa; color: #222; }}
section {{ background: #fff; border: 1px solid #ddd; border-radius: 6px; padding: 12px 18px; margin: 18px 0; }}
h3 {{ margin: 0 0 6px; }} .report {{ background: #eef3fb; padding: 6px 10px; border-radius: 4px; }}
.expected {{ background: #f3f3f3; padding: 6px 10px; border-radius: 4px; color: #555; }} .cut {{ color: #b00; }}
.answer {{ line-height: 1.5; }} .t {{ font-weight: bold; margin-top: 8px; }} .sp {{ height: 8px; }}
.q {{ border-left: 3px solid #ccc; padding-left: 8px; color: #444; }} .li {{ text-indent: -12px; padding-left: 12px; }} nav {{ position: sticky; top: 0; background: #fafafa; padding: 6px 0; border-bottom: 1px solid #ddd; }}
</style></head><body>
<nav>Go to id: <input id="go" size="6" placeholder="F001"> <a href="#files">files</a> · <a href="#photos">photos</a></nav>
{''.join(cards)}
<script>
// just enough markdown to read the answers
function show(text) {{
  const esc = text.replace(/&/g, '&amp;').replace(/</g, '&lt;').replace(/>/g, '&gt;');
  return esc.split('\\n').map(line => {{
    line = line.replace(/\\*\\*(.+?)\\*\\*/g, '<b>$1</b>');
    if (/^\\s*(---|\\*\\*\\*)\\s*$/.test(line)) return '<hr>';
    let m = line.match(/^(#{{1,6}})\\s+(.*)$/);
    if (m) return '<div class="t">' + m[2] + '</div>';
    m = line.match(/^(\\s*)([-*•]|\\d+\\.)\\s+(.*)$/);
    if (m) return '<div class="li" style="margin-left:' + (m[1].length * 8 + 14) + 'px">' + (/\\d/.test(m[2]) ? m[2] : '•') + ' ' + m[3] + '</div>';
    if (/^&gt;/.test(line)) return '<div class="q">' + line.replace(/^&gt;\\s?/, '') + '</div>';
    return line.trim() === '' ? '<div class="sp"></div>' : '<div>' + line + '</div>';
  }}).join('');
}}
document.querySelectorAll('.answer').forEach(d => {{ d.innerHTML = show(d.dataset.md); }});
document.getElementById('go').addEventListener('keydown', e => {{
  if (e.key === 'Enter') location.hash = e.target.value.trim().toUpperCase();
}});
</script></body></html>"""
    path.write_text(page, encoding="utf-8")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    manifest = pd.read_csv(ROOT / "reports" / "manifest.csv", dtype=str).fillna("")
    photos_csv = pd.read_csv(ROOT / "photos" / "photos.csv", dtype=str)[["photo_id", "file"]].rename(columns={"file": "file_photo"})
    keys = []

    files = build(load_answers("full"), manifest, FILE_COLUMNS, "F", keys)

    photo_answers = load_answers("photos")
    photo_answers = photo_answers[(photo_answers.prompt_id == "fr") & photo_answers.report_id.isin(PHOTO_REPORTS)]
    photo_answers = photo_answers.merge(photos_csv, on="photo_id", how="left")
    photos = build(photo_answers, manifest, PHOTO_COLUMNS, "P", keys)

    wb = Workbook()
    rules = wb.active
    rules.title = "rules"
    rules.append(["column", "rule"])
    for name, text in RULES:
        rules.append([name, text])
    rules.column_dimensions["A"].width = 20
    rules.column_dimensions["B"].width = 120
    for row in rules.iter_rows(min_row=1):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    rules["A1"].font = rules["B1"].font = Font(bold=True)

    write_sheet(wb, "files", files, FILE_COLUMNS, condition=False)
    write_sheet(wb, "photos", photos, PHOTO_COLUMNS, condition=True)
    wb.save(OUT / "to_score.xlsx")

    write_html([("files", files, False), ("photos", photos, True)], OUT / "answers.html")
    pd.concat(keys).to_csv(OUT / "key.csv", index=False)
    print(f"{len(files)} answers on files and {len(photos)} on photos written to {OUT}")


if __name__ == "__main__":
    main()
